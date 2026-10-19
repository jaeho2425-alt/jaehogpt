"""Cross-platform pipeline tests. Native Keynote is replaced explicitly, never claimed tested."""
from contextlib import redirect_stdout, redirect_stderr
from io import StringIO
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
import zipfile

import pymupdf as fitz
from pptx import Presentation
from pptx.util import Inches
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'backend'))
import converter as c


def make_pptx(path):
    deck = Presentation()
    for i in range(2):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text = f'NotchConvert Test {i+1}'
    deck.save(path)


def make_pdf(path):
    with fitz.open() as doc:
        for i in range(2):
            page = doc.new_page(width=160, height=100)
            page.insert_text((10, 25), f'NotchConvert Test {i+1}')
            page.draw_rect(fitz.Rect(10, 40, 100, 70), fill=(1, 0, 0))
        doc.save(path)


class PipelineTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.out = self.root / 'results'
        self.pptx = self.root / '발표.pptx'
        make_pptx(self.pptx)
        # Structural placeholders are for routing tests only, not real Keynote/PPT validation.
        self.key = self.root / '발표.key'
        with zipfile.ZipFile(self.key, 'w') as archive:
            archive.writestr('Index/Document.iwa', b'dispatch-test-placeholder')
        self.ppt = self.root / '발표.ppt'
        self.ppt.write_bytes(c.OLE_HEADER + bytes(504))
        self.package = self.root / '패키지.key'
        (self.package / 'Index').mkdir(parents=True)
        (self.package / 'Index/Document.iwa').write_bytes(b'dispatch-test-placeholder')
        self.native = patch.object(c, 'keynote_available', return_value=True)
        self.native.start()
        self.addCleanup(self.native.stop)
    def tearDown(self):
        self.temp.cleanup()

    def fake_export(self, source, fmt, target, work):
        if fmt == 'pdf':
            make_pdf(target)
        else:
            shutil.copyfile(self.pptx, target)

    def test_unsupported_inputs_rejected(self):
        for ext in ['png', 'jpg', 'pdf', 'docx', 'xlsx', 'zip', 'txt']:
            source = self.root / ('input.' + ext)
            source.write_bytes(b'not a presentation')
            with self.subTest(ext=ext), self.assertRaises(c.ConversionError):
                c.inspect([source])

    def test_pptx_copy_preserves_editable_objects_exactly(self):
        with patch.object(c, 'native_export') as export:
            result = c.convert([self.pptx], 'pptx', self.out)
        export.assert_not_called()
        self.assertEqual(Path(result['outputs'][0]).read_bytes(), self.pptx.read_bytes())
        self.assertEqual(len(Presentation(result['outputs'][0]).slides), 2)

    def test_keynote_missing_leaves_only_available_routes(self):
        with patch.object(c, 'keynote_available', return_value=False):
            self.assertEqual(c.inspect([self.key])['formats'], ['zip'])
            self.assertEqual(c.inspect([self.pptx])['formats'], ['pptx', 'zip'])
            self.assertEqual(c.inspect([self.ppt, self.pptx])['formats'], ['zip'])
            with self.assertRaises(c.ConversionError):
                c.convert([self.key], 'pdf', self.out)

    def test_package_zip_preserves_internal_paths(self):
        result = c.convert([self.package], 'zip', self.out)
        with zipfile.ZipFile(result['outputs'][0]) as archive:
            self.assertEqual(archive.read('001/패키지.key/Index/Document.iwa'), b'dispatch-test-placeholder')
            self.assertIsNone(archive.testzip())

    def test_package_symlinks_rejected(self):
        (self.package / 'outside').symlink_to(self.pptx)
        with self.assertRaises(c.ConversionError):
            c.convert([self.package], 'zip', self.out)

    def test_output_inside_package_rejected(self):
        before = sorted(str(p) for p in self.package.rglob('*'))
        with self.assertRaises(c.ConversionError):
            c.convert([self.package], 'zip', self.package / 'results')
        self.assertEqual(sorted(str(p) for p in self.package.rglob('*')), before)

    def test_long_korean_filename(self):
        source = self.root / ('한' * 83 + '.key')
        shutil.copyfile(self.key, source)
        with patch.object(c, 'native_export', side_effect=self.fake_export):
            for fmt in c.FORMATS:
                with self.subTest(fmt=fmt):
                    self.assertTrue(c.convert([source], fmt, self.out)['outputs'])

    def test_corrupt_and_missing_presentations_rejected(self):
        for ext in ['key', 'ppt', 'pptx']:
            path = self.root / ('broken.' + ext)
            path.write_bytes(b'broken')
            with self.subTest(ext=ext), self.assertRaises(c.ConversionError):
                c.inspect([path])
        with self.assertRaises(c.ConversionError):
            c.inspect([self.root / 'missing.pptx'])

    def test_empty_batch_rejected(self):
        with self.assertRaises(c.ConversionError):
            c.convert([], 'zip', self.out)

    def test_first_output_removed_when_second_export_fails(self):
        count = 0
        def export(source, fmt, target, work):
            nonlocal count
            count += 1
            if count == 2:
                raise c.ConversionError('Simulated Keynote failure')
            make_pdf(target)
        with patch.object(c, 'native_export', side_effect=export):
            with self.assertRaises(c.ConversionError):
                c.convert([self.key, self.ppt], 'pdf', self.out)
        self.assertEqual(list(self.out.iterdir()), [])

    def test_repeated_jobs_never_overwrite(self):
        a = c.convert([self.pptx], 'pptx', self.out)
        b = c.convert([self.pptx], 'pptx', self.out)
        self.assertNotEqual(a['folder'], b['folder'])
        self.assertTrue(Path(a['outputs'][0]).is_file())

    def test_concurrent_zip_jobs_are_isolated(self):
        from concurrent.futures import ThreadPoolExecutor
        with ThreadPoolExecutor(max_workers=4) as pool:
            results = list(pool.map(lambda _: c.convert([self.pptx], 'zip', self.out), range(8)))
        self.assertEqual(len({r['folder'] for r in results}), 8)
        self.assertEqual(len(list(self.out.iterdir())), 8)

    def test_read_only_destination(self):
        if hasattr(os, 'geteuid') and os.geteuid() == 0:
            self.skipTest('Requires non-root user')
        self.out.mkdir()
        self.out.chmod(0o500)
        try:
            with self.assertRaises(PermissionError):
                c.convert([self.pptx], 'zip', self.out)
        finally:
            self.out.chmod(0o700)
        self.assertEqual(list(self.out.iterdir()), [])

    def test_slide_limit_cleans_partial_results(self):
        with patch.object(c, 'MAX_PAGES', 1), patch.object(c, 'native_export', side_effect=self.fake_export):
            with self.assertRaises(c.ConversionError):
                c.convert([self.key], 'png', self.out)
        self.assertEqual(list(self.out.iterdir()), [])

    def test_cli_success_and_failure_are_json(self):
        weird = self.root / "한글 ' $(echo nope).pptx"
        shutil.copyfile(self.pptx, weird)
        for source in [weird, self.root / 'missing.pptx']:
            proc = subprocess.run([sys.executable, c.__file__, 'convert', '--format', 'pptx',
                '--output', str(self.out), str(source)], capture_output=True, text=True)
            reply = json.loads(proc.stdout)
            self.assertEqual(reply['ok'], source.exists())
            self.assertEqual(proc.returncode == 0, source.exists())

    def test_diagnostics_cannot_pollute_json(self):
        def noisy(paths):
            print('Library warning')
            return {'formats': ['zip'], 'files': []}
        output, errors = StringIO(), StringIO()
        with patch.object(c, 'inspect', side_effect=noisy), \
             patch.object(sys, 'argv', ['converter.py', 'inspect', str(self.pptx)]), \
             redirect_stdout(output), redirect_stderr(errors):
            c.main()
        self.assertTrue(json.loads(output.getvalue())['ok'])
        self.assertIn('Library warning', errors.getvalue())

    def test_frozen_worker_does_not_leak_library_paths_to_keynote(self):
        values = {'DYLD_LIBRARY_PATH': '/bundled/runtime',
                  'DYLD_FALLBACK_LIBRARY_PATH': '/bundled/fallback',
                  'LD_LIBRARY_PATH': '/bundled/linux',
                  'LD_LIBRARY_PATH_ORIG': '/system/original',
                  'PATH': '/usr/bin:/bin'}
        with patch.dict(os.environ, values, clear=True), patch.object(sys, 'frozen', True, create=True):
            env = c.external_environment()
        self.assertNotIn('DYLD_LIBRARY_PATH', env)
        self.assertNotIn('DYLD_FALLBACK_LIBRARY_PATH', env)
        self.assertEqual(env['LD_LIBRARY_PATH'], '/system/original')
        self.assertEqual(env['PATH'], '/usr/bin:/bin')

    def test_native_uses_private_copy_and_literal_path_arguments(self):
        original = self.key.read_bytes()
        target = self.root / 'output.pdf'
        work = self.root / 'work'
        work.mkdir()
        def execute(args, **kwargs):
            self.assertEqual(args[0], '/usr/bin/osascript')
            self.assertNotEqual(args[2], str(self.key))
            self.assertEqual(Path(args[2]).read_bytes(), original)
            make_pdf(Path(args[3]))
            return subprocess.CompletedProcess(args, 0, '', '')
        with patch.object(c.subprocess, 'run', side_effect=execute):
            c.native_export(self.key, 'pdf', target, work)
        self.assertEqual(self.key.read_bytes(), original)

    def test_native_package_uses_private_copy(self):
        work = self.root / 'work'
        work.mkdir()
        target = self.root / 'output.pdf'
        def execute(args, **kwargs):
            self.assertTrue((Path(args[2]) / 'Index/Document.iwa').is_file())
            self.assertNotEqual(Path(args[2]), self.package)
            make_pdf(Path(args[3]))
            return subprocess.CompletedProcess(args, 0, '', '')
        with patch.object(c.subprocess, 'run', side_effect=execute):
            c.native_export(self.package, 'pdf', target, work)

    def test_permission_timeout_and_missing_output_errors(self):
        for code, message in [('-1743', '제어 권한'), ('-1712', '응답 시간'), ('other', '내보내지')]:
            with self.subTest(code=code), tempfile.TemporaryDirectory() as work:
                with patch.object(c.subprocess, 'run', return_value=subprocess.CompletedProcess([], 1, '', code)):
                    with self.assertRaisesRegex(c.ConversionError, message):
                        c.native_export(self.key, 'pdf', self.root / 'o.pdf', Path(work))
        with tempfile.TemporaryDirectory() as work, patch.object(c.subprocess, 'run', side_effect=subprocess.TimeoutExpired('osascript', 210)):
            with self.assertRaisesRegex(c.ConversionError, '시간'):
                c.native_export(self.key, 'pdf', self.root / 'o.pdf', Path(work))
        with tempfile.TemporaryDirectory() as work, patch.object(c.subprocess, 'run', return_value=subprocess.CompletedProcess([], 0, '', '')):
            with self.assertRaisesRegex(c.ConversionError, '생성하지'):
                c.native_export(self.key, 'pdf', self.root / 'o.pdf', Path(work))


def add_pipeline_case(kind, fmt):
    def test(self):
        source = {'key': self.key, 'ppt': self.ppt, 'pptx': self.pptx, 'key_package': self.package}[kind]
        original = source.read_bytes() if source.is_file() else None
        with patch.object(c, 'native_export', side_effect=self.fake_export) as export:
            result = c.convert([source], fmt, self.out)
        paths = [Path(p) for p in result['outputs']]
        self.assertEqual(len(paths), 2 if fmt in {'jpeg', 'png'} else 1)
        if original is not None:
            self.assertEqual(source.read_bytes(), original)
        for path in paths:
            if fmt in {'jpeg', 'png'}:
                pixmap = fitz.Pixmap(str(path))
                self.assertEqual((pixmap.width, pixmap.height), (320, 200))
                self.assertLess(pixmap.color_topusage()[0], 0.999)
            elif fmt == 'pdf':
                with fitz.open(path) as doc:
                    self.assertEqual(len(doc), 2)
            elif fmt == 'pptx':
                self.assertEqual(len(Presentation(path).slides), 2)
            else:
                with zipfile.ZipFile(path) as archive:
                    self.assertIsNone(archive.testzip())
        if fmt == 'zip' or (kind == 'pptx' and fmt == 'pptx'):
            export.assert_not_called()
        else:
            export.assert_called_once()
    setattr(PipelineTests, f'test_pipeline_{kind}_to_{fmt}', test)

for kind in ['key', 'ppt', 'pptx', 'key_package']:
    for fmt in c.FORMATS:
        add_pipeline_case(kind, fmt)
