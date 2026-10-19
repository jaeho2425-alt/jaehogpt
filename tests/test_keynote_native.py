"""Actual Keynote integration. Opt in on macOS; skipped tests never imply success."""
import hashlib
import base64
import gzip
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
import zipfile

import pymupdf as fitz
from pptx import Presentation
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'backend'))
import converter as c

ROWS = []

def fingerprint(path):
    if path.is_file():
        return hashlib.sha256(path.read_bytes()).hexdigest()
    return [(str(p.relative_to(path)), hashlib.sha256(p.read_bytes()).hexdigest())
            for p in sorted(path.rglob('*')) if p.is_file()]

@unittest.skipUnless(sys.platform == 'darwin' and os.environ.get('NOTCHCONVERT_RUN_KEYNOTE_TESTS') == '1',
                     'Actual Keynote requires an opted-in macOS run')
class NativeKeynoteTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        if not c.keynote_available():
            raise RuntimeError('Keynote를 설치하고 한 번 실행한 뒤 다시 검증해 주세요.')
        cls.temp = tempfile.TemporaryDirectory()
        cls.addClassCleanup(cls.temp.cleanup)
        cls.root = Path(cls.temp.name)
        cls.inputs = {'key': cls.root / 'Native.key', 'pptx': cls.root / 'Native.pptx',
                      'ppt': Path(__file__).parent / 'fixtures/legacy.ppt'}
        if not cls.inputs['ppt'].is_file():
            encoded = Path(__file__).parent / 'fixtures/legacy.ppt.gz.b64.txt'
            cls.inputs['ppt'] = cls.root / 'legacy.ppt'
            cls.inputs['ppt'].write_bytes(gzip.decompress(base64.b64decode(encoded.read_text())))
        result = subprocess.run(['/usr/bin/osascript', str(Path(__file__).with_name('make-keynote-fixture.applescript')),
                                 str(cls.inputs['key']), str(cls.inputs['pptx'])],
                                capture_output=True, text=True, timeout=210)
        if result.returncode:
            raise RuntimeError('Keynote 테스트 파일 생성 실패: ' + result.stderr)
        for path in cls.inputs.values():
            if not path.exists():
                raise RuntimeError(f'테스트 파일이 없습니다: {path}')

    @classmethod
    def tearDownClass(cls):
        report = os.environ.get('NOTCHCONVERT_MATRIX_REPORT')
        if report:
            Path(report).write_text(json.dumps(ROWS, ensure_ascii=False, indent=2))

    def check_native(self, kind, fmt):
        source = self.inputs[kind]
        before = fingerprint(source)
        record = {'input': kind, 'output': fmt, 'status': 'failed', 'engine': 'actual-macOS-Keynote'}
        ROWS.append(record)
        worker = os.environ.get('NOTCHCONVERT_NATIVE_WORKER')
        if worker:
            completed = subprocess.run([worker, 'convert', '--format', fmt, '--output',
                                        str(self.root / 'outputs'), str(source)],
                                       capture_output=True, text=True, timeout=300)
            result = json.loads(completed.stdout)
            self.assertEqual(completed.returncode, 0, completed.stderr + str(result.get('error', '')))
            self.assertTrue(result.get('ok'), result)
            record['runtime'] = 'bundled-worker'
        else:
            result = c.convert([source], fmt, self.root / 'outputs')
            record['runtime'] = 'source-python'
        paths = [Path(p) for p in result['outputs']]
        self.assertEqual(len(paths), 2 if fmt in {'jpeg', 'png'} else 1)
        self.assertEqual(fingerprint(source), before)
        for path in paths:
            if fmt == 'pdf':
                with fitz.open(path) as doc:
                    self.assertEqual(len(doc), 2)
                    self.assertIn('NotchConvert Test', ''.join(p.get_text() for p in doc))
            elif fmt in {'jpeg', 'png'}:
                pixmap = fitz.Pixmap(str(path))
                self.assertGreater(pixmap.width, 0)
                self.assertGreater(pixmap.height, 0)
                self.assertLess(pixmap.color_topusage()[0], 0.999)
            elif fmt == 'pptx':
                deck = Presentation(path)
                self.assertEqual(len(deck.slides), 2)
                text = ''.join(shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
                self.assertIn('NotchConvert Test', text)
            else:
                with zipfile.ZipFile(path) as archive:
                    self.assertIsNone(archive.testzip())
                    if source.is_file():
                        self.assertEqual(archive.read(archive.namelist()[0]), source.read_bytes())
        record['status'] = 'passed'
        record['outputs'] = len(paths)


def add_case(kind, fmt):
    def test(self):
        self.check_native(kind, fmt)
    setattr(NativeKeynoteTests, f'test_native_{kind}_to_{fmt}', test)

for kind in ['key', 'ppt', 'pptx']:
    for fmt in c.FORMATS:
        add_case(kind, fmt)
