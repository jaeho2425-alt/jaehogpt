"""Sibling output, collision handling, and rollback for radial-menu drops."""
import errno
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
import zipfile

import pymupdf as fitz
from pptx import Presentation
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'backend'))
import converter as c

class BesideOriginalTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.source = self.root / '발표.pptx'
        deck = Presentation()
        for _ in range(2):
            deck.slides.add_slide(deck.slide_layouts[6])
        deck.save(self.source)
        self.original = self.source.read_bytes()
        self.native = patch.object(c, 'keynote_available', return_value=True)
        self.native.start()
        self.addCleanup(self.native.stop)
    def tearDown(self):
        self.temp.cleanup()

    def fake_export(self, source, fmt, target, work):
        if fmt == 'pptx':
            shutil.copyfile(self.source, target)
        else:
            with fitz.open() as document:
                for _ in range(2):
                    document.new_page(width=80, height=60).insert_text((10, 20), 'Test')
                document.save(target)

    def test_all_five_outputs_next_to_source(self):
        with patch.object(c, 'native_export', side_effect=self.fake_export):
            for fmt in c.FORMATS:
                with self.subTest(fmt=fmt):
                    result = c.convert_beside_original([self.source], fmt)
                    self.assertEqual(Path(result['folder']), self.root)
                    self.assertEqual(len(result['outputs']), 2 if fmt in {'png', 'jpeg'} else 1)
                    for value in result['outputs']:
                        self.assertEqual(Path(value).parent, self.root)
                        self.assertTrue(Path(value).is_file())
        self.assertEqual(self.source.read_bytes(), self.original)
        self.assertFalse(any(p.name.startswith('.notchconvert') for p in self.root.iterdir()))

    def test_pptx_output_does_not_replace_original(self):
        a = c.convert_beside_original([self.source], 'pptx')
        b = c.convert_beside_original([self.source], 'pptx')
        self.assertEqual(Path(a['outputs'][0]).name, '발표 (2).pptx')
        self.assertEqual(Path(b['outputs'][0]).name, '발표 (3).pptx')
        self.assertEqual(self.source.read_bytes(), self.original)

    def test_existing_pdf_is_untouched(self):
        existing = self.root / '발표.pdf'
        existing.write_bytes(b'keep existing')
        with patch.object(c, 'native_export', side_effect=self.fake_export):
            result = c.convert_beside_original([self.source], 'pdf')
        self.assertEqual(Path(result['outputs'][0]).name, '발표 (2).pdf')
        self.assertEqual(existing.read_bytes(), b'keep existing')

    def test_sources_in_different_folders_stay_in_their_own_folder(self):
        parent = self.root / '다른 폴더'
        parent.mkdir()
        second = parent / self.source.name
        second.write_bytes(self.original)
        with patch.object(c, 'native_export', side_effect=self.fake_export):
            result = c.convert_beside_original([self.source, second], 'png')
        self.assertEqual(len(result['outputs']), 4)
        self.assertEqual({Path(p).parent for p in result['outputs']}, {self.root, parent})
        self.assertEqual(second.read_bytes(), self.original)

    def test_zip_grouped_by_original_folder(self):
        other = self.root / 'subfolder'
        other.mkdir()
        second = other / '다른.pptx'
        second.write_bytes(self.original)
        result = c.convert_beside_original([self.source, second], 'zip')
        self.assertEqual(len(result['outputs']), 2)
        for value in result['outputs']:
            with zipfile.ZipFile(value) as archive:
                self.assertIsNone(archive.testzip())
                self.assertEqual(archive.read(archive.namelist()[0]), self.original)

    def test_failed_later_conversion_leaves_no_visible_outputs(self):
        other = self.root / 'subfolder'
        other.mkdir()
        second = other / '다른.pptx'
        second.write_bytes(self.original)
        calls = 0
        def export(source, fmt, target, work):
            nonlocal calls
            calls += 1
            if calls == 2:
                raise c.ConversionError('Keynote failure')
            self.fake_export(source, fmt, target, work)
        with patch.object(c, 'native_export', side_effect=export):
            with self.assertRaises(c.ConversionError):
                c.convert_beside_original([self.source, second], 'pdf')
        self.assertEqual(sorted(p.name for p in self.root.iterdir()), ['subfolder', self.source.name])
        self.assertEqual([p.name for p in other.iterdir()], [second.name])

    def test_failed_publication_rolls_back_only_new_outputs(self):
        previous = self.root / '발표-001.png'
        previous.write_bytes(b'previous')
        original = c.publish_unique
        calls = 0
        def publish(staged, desired):
            nonlocal calls
            calls += 1
            if calls == 2:
                raise OSError('Disk full')
            return original(staged, desired)
        with patch.object(c, 'native_export', side_effect=self.fake_export), patch.object(c, 'publish_unique', side_effect=publish):
            with self.assertRaises(OSError):
                c.convert_beside_original([self.source], 'png')
        self.assertEqual(previous.read_bytes(), b'previous')
        self.assertEqual(sorted(p.name for p in self.root.iterdir()), sorted([self.source.name, previous.name]))

    def test_no_hardlink_filesystem_uses_exclusive_copy(self):
        with patch.object(c.os, 'link', side_effect=OSError(errno.EOPNOTSUPP, 'Not supported')):
            result = c.convert_beside_original([self.source], 'pptx')
        self.assertEqual(Path(result['outputs'][0]).read_bytes(), self.original)
        self.assertEqual(self.source.read_bytes(), self.original)

    def test_read_only_original_folder_fails_cleanly(self):
        if hasattr(os, 'geteuid') and os.geteuid() == 0:
            self.skipTest('Requires non-root user')
        self.root.chmod(0o500)
        try:
            with self.assertRaises(PermissionError):
                c.convert_beside_original([self.source], 'pptx')
        finally:
            self.root.chmod(0o700)
        self.assertEqual([p.name for p in self.root.iterdir()], [self.source.name])

    def test_concurrent_sibling_copies_have_unique_names(self):
        from concurrent.futures import ThreadPoolExecutor
        with ThreadPoolExecutor(max_workers=4) as pool:
            results = list(pool.map(lambda _: c.convert_beside_original([self.source], 'pptx'), range(8)))
        outputs = [r['outputs'][0] for r in results]
        self.assertEqual(len(set(outputs)), 8)
        self.assertTrue(all(Path(p).read_bytes() == self.original for p in outputs))
        self.assertEqual(self.source.read_bytes(), self.original)

    def test_empty_and_unsupported_batch_are_rejected(self):
        with self.assertRaises(c.ConversionError):
            c.convert_beside_original([], 'png')
        unsupported = self.root / 'photo.png'
        unsupported.write_bytes(b'bad')
        with self.assertRaises(c.ConversionError):
            c.convert_beside_original([unsupported], 'png')

    def test_cli_sibling_mode_keeps_json_and_original(self):
        result = subprocess.run([sys.executable, c.__file__, 'convert', '--format', 'pptx',
                                 '--beside-original', str(self.source)], capture_output=True, text=True)
        self.assertEqual(result.returncode, 0, result.stderr)
        reply = json.loads(result.stdout)
        self.assertTrue(reply['ok'])
        self.assertNotEqual(Path(reply['outputs'][0]), self.source)
        self.assertEqual(self.source.read_bytes(), self.original)
