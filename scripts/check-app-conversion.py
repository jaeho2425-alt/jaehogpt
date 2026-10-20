"""Check the installed GUI's real model/worker connection using a Korean PPTX."""
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import time
import zipfile

from pptx import Presentation


def read_result(report):
    result = json.loads(report.read_text())
    if result['phase'] != 'passed':
        raise RuntimeError(json.dumps(result, ensure_ascii=False))
    return result


def verify(source, result):
    original = source.read_bytes()
    for key in ('pptx', 'nearby'):
        assert Path(result[key]).read_bytes() == original, key
    with zipfile.ZipFile(result['zip']) as archive:
        assert archive.testzip() is None
        assert archive.read(archive.namelist()[0]) == original


def main():
    if sys.platform != 'darwin':
        raise SystemExit('This check requires a macOS graphical session')
    app = Path(sys.argv[1]).resolve(strict=True)
    with tempfile.TemporaryDirectory(prefix='notchconvert GUI ') as temp:
        root = Path(temp).resolve()
        source = root / "한글 ' 제안서.pptx"
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[0])
        slide.shapes.title.text = '변환 연결 검사'
        presentation.save(source)
        try:
            report = root / 'gui.json'
            arguments = ['--check-app-conversion', str(source), str(root / '변환 결과'), str(report)]
            subprocess.run(['/usr/bin/open', '-n', str(app), '--args', *arguments], check=True)
            deadline = time.monotonic() + 90
            while not report.exists():
                if time.monotonic() >= deadline:
                    raise RuntimeError('GUI conversion report missing')
                time.sleep(0.1)
            verify(source, read_result(report))
            print('Installed GUI: inspection, PPTX, ZIP, radial conversion path, original preservation, error replies passed', flush=True)
            # Repeat with no shell/Python setup and an explicitly ASCII stdio locale.
            report = root / 'minimal.json'
            arguments[-1] = str(report)
            environment = {key: os.environ[key] for key in ('HOME', 'TMPDIR', '__CF_USER_TEXT_ENCODING') if key in os.environ}
            environment.update(PATH='/usr/bin:/bin:/usr/sbin:/sbin', LANG='C', LC_ALL='C',
                               PYTHONUTF8='0', PYTHONCOERCECLOCALE='0', PYTHONIOENCODING='ascii')
            process = subprocess.run([str(app / 'Contents/MacOS/NotchConvert'), *arguments],
                                     cwd='/', env=environment, stdin=subprocess.DEVNULL,
                                     capture_output=True, timeout=90)
            if not report.exists():
                raise RuntimeError('Minimal environment GUI failed: ' + process.stderr.decode('utf-8', errors='replace')[-2000:])
            verify(source, read_result(report))
            print('Installed GUI with minimal environment and ASCII locale: passed', flush=True)
        finally:
            subprocess.run(['/usr/bin/pkill', '-x', 'NotchConvert'], check=False)


if __name__ == '__main__':
    main()
